"""
fayl/oqi.py — Word (.docx) va PDF fayldan matn oladi (jadval ichini ham).
"""
import os


def matn_ol(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return _docx(path)
    if ext == ".pdf":
        return _pdf(path)
    raise ValueError("Faqat .docx yoki .pdf qabul qilinadi.")


_QIZIL_CHEGARA = 100  # shu qiymatdan past qizil (R yuqori, G/B past) rangni "belgilangan" deb hisoblaymiz


def _qizilmi(run):
    try:
        rgb = run.font.color.rgb  # RGBColor — 3 baytli (r, g, b), masalan RGBColor(0xFF,0,0)
    except AttributeError:
        return False
    if rgb is None:
        return False
    r, g, b = rgb[0], rgb[1], rgb[2]
    return r > 150 and g < _QIZIL_CHEGARA and b < _QIZIL_CHEGARA


def _paragraf_matni(p):
    """Paragraf matnini qaytaradi; agar biror run qizil rangda bo'lsa,
    uni [TOʻGʻRI] bilan belgilaydi (o'qituvchilar ko'pincha to'g'ri
    javobni qizil rangda belgilaydi, matnda alohida belgi bo'lmaydi)."""
    matn = p.text
    if not matn.strip():
        return matn
    if any(_qizilmi(r) for r in p.runs if r.text.strip()):
        return f"[TO'G'RI] {matn}"
    return matn


def _docx(path):
    """Bo'sh paragraflar HAM bitta bo'sh qator sifatida saqlanadi (ketma-ket
    bo'shlar bittaga qisqartiriladi) — bu raqamlanmagan hujjatlarda savol
    bloklarini bo'sh qator chegarasi bo'yicha aniqlash imkonini beradi."""
    from docx import Document
    doc = Document(path)
    qatorlar = []
    oldingi_bosh = True  # boshida bo'sh qator kerak emas
    for p in doc.paragraphs:
        if p.text.strip():
            qatorlar.append(_paragraf_matni(p))
            oldingi_bosh = False
        elif not oldingi_bosh:
            qatorlar.append("")
            oldingi_bosh = True
    # jadval ichidagi matn ham. Savol+variantlar bitta katakda (ko'p
    # paragraf sifatida) bo'lishi mumkin — shuning uchun har katakning
    # o'z paragraflarini (rang-belgisi bilan birga, xuddi oddiy matndagi
    # kabi) alohida qatorlar deb olamiz, aks holda ko'p qatorli katak
    # qo'shni katak bilan bitta qatorga qo'shilib, tuzilishi buzilib
    # qolardi VA rang orqali belgilangan to'g'ri javob yo'qolib qolardi.
    # Har jadval QATORI orasiga bo'sh qator qo'shamiz (odatda 1 qator =
    # 1 savol bo'ladi) — shu orqali raqamlanmagan jadvalli testlar ham
    # bo'sh-qator-chegarasi bo'yicha ajratilishi mumkin bo'ladi.
    if qatorlar and qatorlar[-1] != "":
        qatorlar.append("")  # paragraf va jadval matni bir blokka qo'shilib qolmasin
    for t in doc.tables:
        for row in t.rows:
            katak_qatorlari = []
            for c in row.cells:
                c_qatorlar = [_paragraf_matni(p) for p in c.paragraphs if p.text.strip()]
                if c_qatorlar:
                    katak_qatorlari.append(c_qatorlar)
            if not katak_qatorlari:
                continue
            if any(len(k) > 1 for k in katak_qatorlari):
                for k in katak_qatorlari:
                    qatorlar.extend(k)
            else:
                qatorlar.append("  ".join(k[0] for k in katak_qatorlari))
            qatorlar.append("")
    return "\n".join(qatorlar)


def _pdf_soz_qizilmi(soz):
    rang = soz.get("non_stroking_color")
    if not rang or len(rang) < 3:
        return False
    r, g, b = rang[0], rang[1], rang[2]  # pdfplumber: 0..1 oralig'ida
    return r > 0.55 and g < 0.4 and b < 0.4


def _pdf_qatorlar(page):
    """Sahifa matnini qatorlarga guruhlab, ko'pchilik so'zi qizil rangda
    bo'lgan qatorni [TO'G'RI] bilan belgilaydi (docx dagi kabi)."""
    try:
        sozlar = page.extract_words(extra_attrs=["non_stroking_color"])
    except Exception:
        sozlar = page.extract_words()
    guruh = {}
    for s in sozlar:
        guruh.setdefault(round(s["top"]), []).append(s)
    natija = []
    for top in sorted(guruh):
        qator_sozlari = guruh[top]
        matn = " ".join(s["text"] for s in qator_sozlari)
        if not matn.strip():
            continue
        qizil = sum(1 for s in qator_sozlari if _pdf_soz_qizilmi(s))
        if qizil and qizil >= len(qator_sozlari) / 2:
            matn = f"[TO'G'RI] {matn}"
        natija.append(matn)
    return natija


def _pdf(path):
    import pdfplumber
    qatorlar = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            qatorlar.extend(_pdf_qatorlar(page))
            # jadvallar (katak ichida ko'p qatorli matn bo'lsa, qo'shni
            # katak bilan qo'shilib ketmasligi uchun alohida qatorlarga
            # bo'lamiz; har qator orasiga bo'sh qator — odatda 1 jadval
            # qatori = 1 savol bo'ladi)
            for tbl in (page.extract_tables() or []):
                for row in tbl:
                    hujayralar = [str(c).strip() for c in row if c and str(c).strip()]
                    if not hujayralar:
                        continue
                    if any("\n" in h for h in hujayralar):
                        for h in hujayralar:
                            qatorlar.extend(h.split("\n"))
                    else:
                        qatorlar.append("  ".join(hujayralar))
                    qatorlar.append("")
    return "\n".join(qatorlar)
